from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt


def set_cell(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    run.font.name = "Calibri"


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, bold=True)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            set_cell(table.rows[row_index + 1].cells[col_index], value)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("BoardVerse — Báo cáo kiểm thử API theo role", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        "Môi trường: https://boardverse-server.onrender.com\n"
        "Ngày kiểm thử: 12/08/2026\n"
        "Phạm vi: Smoke test GET (+ một số mutation Admin)\n"
        "Roles: Admin, CafeStaff"
    ).font.size = Pt(11)

    doc.add_paragraph(
        "Tài liệu ghi lại các chức năng/API đã kiểm thử thành công (OK) và chưa đạt / bị chặn / còn hạn chế "
        "đối với hai role Admin và Staff (CafeStaff)."
    )

    # ===================== ADMIN =====================
    doc.add_heading("1. Role Admin", level=1)
    doc.add_paragraph(
        "Tài khoản test: admin (role Admin). Token JWT hợp lệ. "
        "Kiểm thử trên Swagger/live API và đối chiếu FE Admin Portal."
    )

    doc.add_heading("1.1. Đã test OK", level=2)
    add_table(
        doc,
        ["Nhóm chức năng", "Endpoint / Hành vi", "Kết quả"],
        [
            ["Auth", "Login Admin → JWT Bearer", "OK"],
            ["Moderation", "GET /api/v1/admin/karma-logs", "200 OK"],
            ["Moderation", "GET /api/v1/admin/users/alerts (Karma < 50)", "200 OK"],
            ["Moderation", "POST /api/v1/admin/users/{id}/punish (Warning)", "200 OK"],
            ["Moderation", "POST /api/v1/admin/users/{id}/adjust-karma", "200 OK (45→46)"],
            ["Reports", "GET /api/v1/admin/reports/overview", "200 OK"],
            ["Reports", "GET /api/v1/admin/reports/lobby-failures", "200 OK"],
            ["Reports", "GET /api/v1/admin/reports/deposits", "200 OK"],
            ["Reports", "GET /api/v1/admin/reports/cafe-performance", "200 OK"],
            ["Tournament", "GET /api/v1/admin/tournaments", "200 OK"],
            ["Tournament", "GET /api/v1/admin/tournaments/{id}", "200 OK"],
            ["Tournament", "GET .../tournaments/{id}/participants", "200 OK"],
            ["Tournament", "POST .../close-registration", "200 OK"],
            ["Wallet", "GET /api/v1/admin/wallet", "200 OK"],
            ["Wallet", "GET /api/v1/admin/wallet/refund-requests", "200 OK (list rỗng)"],
            ["Configs", "GET /api/v1/admin/configs", "200 OK"],
            ["Categories", "GET /api/v1/admin/categories", "200 OK"],
            ["Cafes", "GET /api/v1/admin/cafes (+ detail)", "200 OK"],
            ["Master catalog", "GET .../master-games/{id}/components", "200 OK"],
            ["Master catalog", "GET .../master-games/{id}/categories", "200 OK"],
            ["SePay", "GET /api/sepay-accounts", "200 OK"],
            ["SePay", "GET /api/sepay-accounts/master", "200 OK"],
            ["User Mgmt", "GET /api/UserManagement/users", "200 OK"],
            ["User Mgmt", "PUT /api/UserManagement/users/{id}/role", "200 OK"],
        ],
    )

    doc.add_heading("1.2. Chưa được / còn hạn chế", level=2)
    add_table(
        doc,
        ["Nhóm", "Vấn đề", "Ghi chú"],
        [
            [
                "SePay path cũ",
                "GET /api/admin/sepay-accounts",
                "404 — path đúng là /api/sepay-accounts (FE đã dùng path đúng)",
            ],
            [
                "Tournament path docs cũ",
                "POST .../registration/open",
                "404 — Swagger live dùng open-registration / close-registration (FE đã sửa)",
            ],
            [
                "Tournament check-in",
                "Check-in participant trên Admin",
                "Admin controller không có check-in; POS endpoint thường 403 với Admin",
            ],
            [
                "Tournament reopen",
                "Reopen registration sau khi đóng",
                "Admin không có reopen; POS reopen-registration → 403 với Admin",
            ],
            [
                "Settlement override",
                "POST .../settlements/{id}/override",
                "Route OK nhưng 404 với id giả; chưa có list Failed để thao tác thực tế",
            ],
            [
                "Refund resolve",
                "POST .../refund-requests/{id}/resolve",
                "Thiếu Idempotency-Key → 400 (FE đã bổ sung header); chưa có request Pending thật để verify Approve/Reject",
            ],
            [
                "Tournament create body",
                "Field prizePoolBvc",
                "Swagger create/update không có prizePool — FE gửi sẽ bị bỏ qua",
            ],
            [
                "BE quirk",
                "punish actionType=Warning",
                "HTTP 200 nhưng response trả accountStatus=Suspended (hành vi BE lạ)",
            ],
            [
                "Data side-effect",
                "Giải Splendor Tourment",
                "Sau smoke close-registration đang ở trạng thái RegistrationClosed",
            ],
            [
                "FE reports (đã vá)",
                "Shape cafe-performance / deposits / lobby",
                "Live trả cafes / flat totals / id+title — FE mapper đã chỉnh theo live",
            ],
        ],
    )

    doc.add_heading("1.3. Module FE Admin đã bổ sung (liên quan test)", level=2)
    for item in [
        "Karma logs + alerts + punish + adjust-karma",
        "Tournaments list/detail/lifecycle",
        "Reports (overview / lobby-failures / deposits / cafe-performance)",
        "Settlement override form",
        "SePay accounts",
        "Wallet refund-requests",
        "Master game: DELETE component, categories, metadata, thumbnail",
        "User: đổi role riêng PUT .../users/{id}/role",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # ===================== STAFF =====================
    doc.add_heading("2. Role Staff (CafeStaff)", level=1)
    doc.add_paragraph(
        "Tài khoản test: staff@gmail.com (role CafeStaff). "
        "Quán được gán: Boss cafe (a1aae9db-4f1b-44af-ac86-6038d085df94)."
    )

    doc.add_heading("2.1. Đã test OK", level=2)
    add_table(
        doc,
        ["Nhóm chức năng", "Endpoint / Hành vi", "Kết quả"],
        [
            ["Auth", "Login Staff → JWT CafeStaff", "OK"],
            ["Staff core", "GET /api/staff/my-cafes", "200 OK — trả Boss cafe"],
            ["Cafe", "GET /api/cafes/{cafeId}", "200 OK"],
            ["POS", "GET .../pos/tables", "200 OK"],
            ["POS", "GET .../pos/boxes", "200 OK"],
            ["POS", "GET .../pos/sessions/active", "200 OK (có phiên)"],
            ["POS", "GET .../pos/sessions/unpaid", "200 OK"],
            ["POS", "GET .../pos/sessions/paid", "200 OK"],
            ["Settlement", "GET .../settlements/pending", "200 OK"],
            ["Inventory", "GET .../inventory", "200 OK"],
        ],
    )

    doc.add_heading("2.2. Chưa được / bị chặn đúng quyền / hạn chế", level=2)
    add_table(
        doc,
        ["Nhóm", "Endpoint / Hành vi", "Kết quả / Ghi chú"],
        [
            ["Manager-only", "GET /api/manager/my-cafes", "403 — đúng (Staff không phải Manager)"],
            ["Admin-only", "GET /api/v1/admin/reports/overview", "403 — đúng"],
            ["Admin-only", "GET /api/UserManagement/users", "403 — đúng"],
            ["Manager-only", "GET /api/cafes/{id}/staff", "403 — quản lý staff thuộc Manager"],
            [
                "Tournament POS",
                "GET /api/v1/pos/tournaments/cafes/{id}",
                "403 — Tournament POS dành Manager",
            ],
            ["Tournament POS", "GET .../pos/tournaments/cafes/{id}/active", "403"],
            [
                "Shifts",
                "GET /api/shifts/current|history",
                "404 — BE resolve cafeId = 0000… (thiếu context quán)",
            ],
            [
                "Mutation POS",
                "check-in / end session / pay / …",
                "Chưa chạy smoke mutation trong đợt test này",
            ],
        ],
    )

    doc.add_heading("2.3. Kết luận Staff", level=2)
    doc.add_paragraph(
        "Staff dùng được luồng vận hành quán: my-cafes → POS (tables/boxes/sessions) + inventory + settlements pending. "
        "Không được vào Admin, UserManagement, quản lý staff list, và Tournament POS. "
        "FE Staff nên dựa /api/staff/my-cafes rồi gọi POS/inventory theo cafeId; không gọi tournament POS nếu chỉ role CafeStaff."
    )

    # ===================== SUMMARY =====================
    doc.add_heading("3. Tổng hợp nhanh", level=1)
    add_table(
        doc,
        ["Role", "OK chính", "Chưa OK / hạn chế chính"],
        [
            [
                "Admin",
                "Moderation, Reports, Tournament list/detail/lifecycle, Wallet list/refunds list, Configs, Categories, Cafes, Master catalog, SePay, đổi role",
                "Check-in tournament (không có trên Admin), reopen registration, settlement list Failed, refund resolve chưa E2E Pending, prizePool không có trên create DTO",
            ],
            [
                "Staff",
                "my-cafes, cafe detail, POS tables/boxes/sessions, inventory, settlements pending",
                "403 đúng với Admin/Manager/tournament POS; shifts 404; chưa smoke mutation POS",
            ],
        ],
    )

    doc.add_heading("4. Khuyến nghị tiếp theo", level=1)
    for item in [
        "Admin FE: ẩn/disable nút check-in participant (hoặc chuyển sang luồng Manager POS).",
        "Admin FE: bỏ/ẩn prizePool trên form create tournament nếu BE không nhận.",
        "Admin: bổ sung UI list settlement Failed trước khi override.",
        "Tạo 1 refund-request Pending thật để verify resolve + Idempotency-Key.",
        "Staff: kiểm tra thêm mutation POS (check-in booking, end session, pay) với cafe Boss.",
        "BE: làm rõ shifts API (cần cafeId?) và hành vi punish Warning vs Suspended.",
        "Manager cần reopen registration cho giải Splendor nếu muốn mở lại ĐK.",
    ]:
        doc.add_paragraph(item, style="List Number")

    footer = doc.add_paragraph()
    run = footer.add_run(
        f"\n— Hết báo cáo —\nSinh tự động từ kết quả smoke test Cursor Agent · "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    run.italic = True

    out_dir = Path(r"D:\CAPSTONE PROJECT KÌ 9\boardverse")
    out_path = out_dir / "Bao-cao-test-API-Admin-Staff.docx"
    doc.save(str(out_path))
    print(f"SAVED {out_path}")

    downloads = Path(r"C:\Users\PC\Downloads") / "Bao-cao-test-API-Admin-Staff.docx"
    try:
        doc.save(str(downloads))
        print(f"SAVED {downloads}")
    except Exception as exc:  # noqa: BLE001
        print(f"DOWNLOADS_FAIL {exc}")


if __name__ == "__main__":
    main()
